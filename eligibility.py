import re
from docx import Document
from jobs import JOB_DATABASE


def extract_resume_text(file_path):
    try:
        document = Document(file_path)
        text = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text.strip())

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text.strip())

        return "\n".join(text)

    except Exception as exc:
        raise RuntimeError(f"Could not read resume: {exc}")


def extract_years_of_experience(resume_text):
    text = resume_text.lower()

    patterns = [
        r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s*(?:of\s*)?experience",
        r"experience\s*[:\-]?\s*(\d+(?:\.\d+)?)\s*\+?\s*years?",
        r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s*(?:in|working|worked)"
    ]

    values = []

    for pattern in patterns:
        for match in re.findall(pattern, text):
            try:
                values.append(float(match))
            except ValueError:
                pass

    if values:
        return max(values)

    if any(x in text for x in [
        "fresher",
        "fresh graduate",
        "recent graduate",
        "student",
        "undergraduate",
        "entry level"
    ]):
        return 0.0

    return 0.0


def parse_experience_requirement(experience_string):
    text = experience_string.lower()

    match = re.search(r"(\d+(?:\.\d+)?)\s*\+", text)
    if match:
        return float(match.group(1))

    match = re.search(r"(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)", text)
    if match:
        return float(match.group(1))

    match = re.search(r"(\d+(?:\.\d+)?)", text)
    if match:
        return float(match.group(1))

    return 0.0


def normalize_text(text):
    text = text.lower()

    replacements = {
        "c plus plus": "c++",
        "cplusplus": "c++",
        "c sharp": "c#",
        "javascript": "javascript",
        "js": "javascript",
        "typescript": "typescript",
        "ts": "typescript",
        "postgres": "postgresql",
        "power bi": "powerbi",
        "machine-learning": "machine learning",
        "deep-learning": "deep learning",
        "restful api": "rest api",
        "restful apis": "rest apis"
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text


def skill_present(skill, resume_text):
    skill = normalize_text(skill)
    resume = normalize_text(resume_text)

    if skill in resume:
        return True

    aliases = {
        "c/c++": ["c++", "c", "cpp"],
        "python/java": ["python", "java"],
        "aws/gcp": ["aws", "gcp"],
        "aws/azure": ["aws", "azure"],
        "tableau/powerbi": ["tableau", "powerbi"],
        "x-ray/mri certification": ["x-ray", "mri", "radiology"],
        "pharm.d or b.pharm": ["pharm.d", "b.pharm"],
        "mbbs degree": ["mbbs"],
        "master's/ph.d. in field": ["master", "ph.d", "phd"]
    }

    for alias in aliases:
        if skill == alias:
            return any(x in resume for x in aliases[alias])

    words = [w for w in skill.split() if len(w) > 2]

    if len(words) > 1:
        return all(word in resume for word in words)

    return False


def get_categories():
    return list(JOB_DATABASE.keys())


def get_jobs(category):
    if category not in JOB_DATABASE:
        return []

    return [job["title"] for job in JOB_DATABASE[category]]


def get_job(category, job_title):
    if category not in JOB_DATABASE:
        return None

    for job in JOB_DATABASE[category]:
        if job["title"].lower() == job_title.lower():
            return job

    return None


def check_eligibility(category, job_title, resume_path):
    job = get_job(category, job_title)

    if not job:
        return {
            "success": False,
            "error": "Job not found."
        }

    resume_text = extract_resume_text(resume_path)

    if not resume_text.strip():
        return {
            "success": False,
            "error": "Resume is empty."
        }

    required_skills = job.get("requirements", [])

    matched_requirements = []
    missing_requirements = []

    for requirement in required_skills:
        if skill_present(requirement, resume_text):
            matched_requirements.append(requirement)
        else:
            missing_requirements.append(requirement)

    candidate_experience = extract_years_of_experience(resume_text)

    required_experience = parse_experience_requirement(
        job.get("experience", "0 years")
    )

    experience_gap = max(
        0,
        required_experience - candidate_experience
    )

    experience_met = candidate_experience >= required_experience

    total_requirements = len(required_skills)

    if total_requirements:
        skill_score = len(matched_requirements) / total_requirements
    else:
        skill_score = 1.0

    if required_experience > 0:
        experience_score = min(
            candidate_experience / required_experience,
            1.0
        )
    else:
        experience_score = 1.0

    score = (
        skill_score * 70 +
        experience_score * 30
    )

    if not missing_requirements and experience_met:
        status = "Eligible"
    elif score >= 40:
        status = "Partially Eligible"
    else:
        status = "Not Eligible"

    return {
        "success": True,
        "category": category,
        "job": job["title"],
        "status": status,
        "score": round(score, 1),
        "experience_required": job["experience"],
        "required_experience_years": required_experience,
        "candidate_experience_years": candidate_experience,
        "experience_gap_years": experience_gap,
        "matched_requirements": matched_requirements,
        "missing_requirements": missing_requirements,
        "remaining_requirements": missing_requirements,
        "matched_count": len(matched_requirements),
        "missing_count": len(missing_requirements),
        "total_requirements": total_requirements
    }